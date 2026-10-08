#!/usr/bin/env python3
"""Turn Denslow's Three Bears (Project Gutenberg #19772) into book pages and a slideshow,
like the slideshows of picture-book pages that bookTranslate was written for.

usage: python demo/make_pages.py BOOK_FOLDER
BOOK_FOLDER holds the book's HTML file from Project Gutenberg (book.html) and its images/ folder.
Writes BOOK_FOLDER/pages/pageNN.jpg and BOOK_FOLDER/slides/Denslows Three Bears.pptx.
Needs Pillow and python-pptx.
"""
import datetime, itertools, math, os, re, sys, textwrap
from html.parser import HTMLParser
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu

SKIP = {"images/i001a.jpg"}            # dedication page
END = "Denslow's Picture Books for Children"   # the publisher's list of other books starts here
LIMIT = 130                            # words per text page
FONT = "/System/Library/Fonts/Supplemental/Georgia.ttf"

class Book(HTMLParser):
  """Pictures and paragraphs, in page order."""
  def __init__(self):
    super().__init__()
    self.items, self.text, self.inBody, self.skip = [], "", False, 0
  def flush(self):
    text = re.sub(r"\s+", " ", self.text).strip()
    if text:
      self.items.append(("text", text))
    self.text = ""
  def handle_starttag(self, tag, attrs):
    if tag == "body":
      self.inBody = True
    if tag in ("style", "script"):
      self.skip += 1
    if tag in ("p", "h1", "h2", "h3", "div", "br"):
      self.flush()
    if tag == "img" and self.inBody:
      self.flush()
      self.items.append(("img", dict(attrs).get("src")))
  def handle_endtag(self, tag):
    if tag in ("style", "script"):
      self.skip -= 1
    if tag in ("p", "h1", "h2", "h3", "div"):
      self.flush()
  def handle_data(self, data):
    if self.inBody and not self.skip:
      self.text += data

def split_run(paras):
  """Split paragraphs that follow each other into the fewest pages with the most even lengths."""
  words = [len(p.split()) for p in paras]
  pages = math.ceil(sum(words) / LIMIT)
  best = None
  for cuts in itertools.combinations(range(1, len(paras)), pages - 1):
    bounds = (0, *cuts, len(paras))
    longest = max(sum(words[a:b]) for a, b in zip(bounds, bounds[1:]))
    if best is None or longest < best[0]:
      best = (longest, bounds)
  return [paras[a:b] for a, b in zip(best[1], best[1][1:])]

folder = sys.argv[1]
book = Book()
book.feed(open(os.path.join(folder, "book.html"), encoding="utf-8").read())
book.flush()
items = book.items
start = next(i for i, item in enumerate(items) if item[0] == "img")
end = next(i for i, item in enumerate(items) if item[0] == "text" and item[1].startswith(END))

pages, run = [], []
for item in items[start:end] + [("img", None)]:
  if item[0] == "img":
    pages += [("text", chunk) for chunk in split_run(run)] if run else []
    run = []
    if item[1] and item[1] not in SKIP:
      pages.append(("img", item[1]))
  else:
    run.append(item[1])

W, H, M = 760, 960, 60
serif = ImageFont.truetype(FONT, 28)
os.makedirs(os.path.join(folder, "pages"), exist_ok=True)
files = []
for n, page in enumerate(pages, start=1):
  out = os.path.join(folder, "pages", f"page{n:02d}.jpg")
  if page[0] == "img":
    Image.open(os.path.join(folder, page[1])).convert("RGB").save(out, quality=90)
  else:
    img = Image.new("RGB", (W, H), (250, 245, 232))
    draw = ImageDraw.Draw(img)
    y = M
    for para in page[1]:
      for line in textwrap.wrap(para, width=44):
        draw.text((M, y), line, font=serif, fill=(40, 34, 28))
        y += 40
      y += 20
    assert y < H - M + 20, f"text does not fit on page {n}"
    img.save(out, quality=90)
  files.append(out)

slides = Presentation()
slides.slide_width, slides.slide_height = Emu(9144000), Emu(6858000)   # 4:3
props = slides.core_properties                                         # not python-pptx's template details
props.title, props.last_modified_by, props.comments = "Denslow's Three Bears", "", ""
props.created = props.modified = datetime.datetime(2026, 10, 8)
for f in files:
  slide = slides.slides.add_slide(slides.slide_layouts[6])             # blank
  w, h = Image.open(f).size
  scale = min(slides.slide_width / w, slides.slide_height / h)
  slide.shapes.add_picture(f, int(slides.slide_width - w * scale) // 2, int(slides.slide_height - h * scale) // 2,
                           int(w * scale), int(h * scale))
os.makedirs(os.path.join(folder, "slides"), exist_ok=True)
slides.save(os.path.join(folder, "slides", "Denslows Three Bears.pptx"))
print(len(files), "pages")
